import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.chart import BarChart, LineChart, Reference

# Load data
df = pd.read_excel("data/sales_data.xlsx")

df["Order_Date"] = pd.to_datetime(df["Order_Date"])

# KPIs
total_sales = df["Sales"].sum()
total_profit = df["Profit"].sum()
total_orders = df["Order_ID"].nunique()
total_quantity = df["Quantity"].sum()
average_order_value = total_sales / total_orders

# Create Excel report
with pd.ExcelWriter("output/Ecommerce_Sales_Report.xlsx", engine="openpyxl") as writer:

    # Raw data
    df.to_excel(writer, sheet_name="Sales Data", index=False)

    # KPI Summary
    kpi = pd.DataFrame({
        "Metric": [
            "Total Sales",
            "Total Profit",
            "Total Orders",
            "Total Quantity",
            "Average Order Value"
        ],
        "Value": [
            total_sales,
            total_profit,
            total_orders,
            total_quantity,
            average_order_value
        ]
    })

    kpi.to_excel(writer, sheet_name="Dashboard", index=False)

    # Category analysis
    category = (
        df.groupby("Category")["Sales"]
        .sum()
        .sort_values(ascending=False)
        .reset_index()
    )

    category.to_excel(
        writer,
        sheet_name="Category Analysis",
        index=False
    )

    # Region analysis
    region = (
        df.groupby("Region")["Sales"]
        .sum()
        .sort_values(ascending=False)
        .reset_index()
    )

    region.to_excel(
        writer,
        sheet_name="Region Analysis",
        index=False
    )

    # Monthly analysis
    df["Month"] = df["Order_Date"].dt.to_period("M").astype(str)

    monthly = (
        df.groupby("Month")["Sales"]
        .sum()
        .reset_index()
    )

    monthly.to_excel(
        writer,
        sheet_name="Monthly Analysis",
        index=False
    )

# Format workbook
wb = load_workbook("output/Ecommerce_Sales_Report.xlsx")

for ws in wb.worksheets:

    # Header formatting
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = max_length + 2

# Add bar chart to Dashboard
dashboard = wb["Dashboard"]

chart = BarChart()
chart.title = "Key Metrics"
chart.y_axis.title = "Value"
chart.x_axis.title = "Metric"

data = Reference(dashboard, min_col=2, min_row=1, max_row=6)
categories = Reference(dashboard, min_col=1, min_row=2, max_row=6)

chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)

dashboard.add_chart(chart, "D2")

wb.save("output/Ecommerce_Sales_Report.xlsx")

print("Excel report created successfully!")
print("File: output/Ecommerce_Sales_Report.xlsx")